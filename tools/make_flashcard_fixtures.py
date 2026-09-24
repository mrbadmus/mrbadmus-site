#!/usr/bin/env python3
"""MRB-351 — build the flashcard-extraction fixture corpus.

Writes tests/fixtures/flashcards/<name>.<ext> plus expected.json, which holds
the question–answer pairs each file is meant to yield. The files are made to
look like what teachers actually hand in: a Rainford-style pptx with the answer
on the NEXT slide, answers in speaker notes, answers in red on the same slide,
a docx table, a docx with an answer section at the end, a text pdf, a scanned
(image-only) pdf, a photographed sheet, a csv, an xlsx keyword list, cloze
sentences, a questions-only list, and a deck carrying a class list that must
never come back out.

Re-run only when the corpus itself changes:
    python3 tools/make_flashcard_fixtures.py
The generated files are committed; CI reads them, it does not rebuild them.
"""
import csv
import io
import json
import os
import random

from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "fixtures", "flashcards")
OUT = os.path.normpath(OUT)
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

# ── the content ─────────────────────────────────────────────────────────────
FORCES_20 = [
    ("What is the unit of force?", "The newton (N)"),
    ("What is weight?", "The force acting on an object due to gravity"),
    ("Write the equation that links weight, mass and gravitational field strength.", "Weight = mass × gravitational field strength (W = mg)"),
    ("What is the gravitational field strength on Earth?", "9.8 N/kg"),
    ("What is a resultant force?", "The single force that has the same effect as all the forces acting on an object"),
    ("What happens to a stationary object when the resultant force on it is zero?", "It stays stationary"),
    ("What is friction?", "A force that opposes motion between two surfaces in contact"),
    ("What is a contact force?", "A force that acts only when objects are touching"),
    ("Give an example of a non-contact force.", "Gravitational, magnetic or electrostatic force"),
    ("What is the difference between mass and weight?", "Mass is the amount of matter and is measured in kg; weight is a force measured in N"),
    ("What is air resistance?", "A frictional force that acts on an object moving through air"),
    ("Write the equation linking work done, force and distance.", "Work done = force × distance (W = Fs)"),
    ("What is the unit of work done?", "The joule (J)"),
    ("What does Hooke's law state?", "The extension of a spring is directly proportional to the force applied, up to the limit of proportionality"),
    ("Write the equation for the force on a spring.", "Force = spring constant × extension (F = ke)"),
    ("What is a vector quantity?", "A quantity that has both magnitude and direction"),
    ("What is a scalar quantity?", "A quantity that has magnitude only"),
    ("Is velocity a scalar or a vector?", "A vector"),
    ("What is terminal velocity?", "The constant velocity reached when the resultant force on a falling object is zero"),
    ("What does Newton's first law state?", "An object stays at rest or moves at constant velocity unless a resultant force acts on it"),
]

CELLS_NOTES = [
    ("What is the function of the nucleus?", "It contains genetic material and controls the activities of the cell"),
    ("What is the function of the cell membrane?", "It controls what enters and leaves the cell"),
    ("Where does aerobic respiration take place?", "In the mitochondria"),
    ("What is the function of ribosomes?", "They are where protein synthesis takes place"),
    ("Name two structures found in plant cells but not animal cells.", "Cell wall and chloroplasts (also a permanent vacuole)"),
    ("What is the cell wall made of in plants?", "Cellulose"),
    ("What happens in chloroplasts?", "Photosynthesis"),
    ("What type of cell has no nucleus: prokaryotic or eukaryotic?", "Prokaryotic"),
]

ACIDS_RED = [
    ("What is the pH of a neutral solution?", "7"),
    ("Which ion makes a solution acidic?", "Hydrogen ions, H+"),
    ("Which ion makes a solution alkaline?", "Hydroxide ions, OH-"),
    ("What is produced when an acid reacts with an alkali?", "A salt and water"),
    ("Name the salt made when hydrochloric acid reacts with sodium hydroxide.", "Sodium chloride"),
    ("What gas is produced when an acid reacts with a metal carbonate?", "Carbon dioxide"),
    ("What is the formula of sulfuric acid?", "H2SO4"),
    ("What colour is universal indicator in a strong acid?", "Red"),
]

ENERGY_TABLE = [
    ("Name the energy store in a stretched spring.", "Elastic potential energy store"),
    ("Name the energy store of a moving object.", "Kinetic energy store"),
    ("What is the unit of energy?", "The joule (J)"),
    ("What does the law of conservation of energy state?", "Energy cannot be created or destroyed, only transferred"),
    ("What is power?", "The rate at which energy is transferred"),
    ("What is the unit of power?", "The watt (W)"),
    ("Write the equation for efficiency.", "Efficiency = useful output energy ÷ total input energy"),
    ("What happens to wasted energy?", "It is dissipated to the surroundings, usually as heat"),
    ("How can unwanted energy transfer be reduced in a machine?", "By lubrication"),
    ("Name a renewable energy resource.", "Wind, solar, tidal, hydroelectric, geothermal or biomass"),
    ("Name a non-renewable energy resource.", "Coal, oil, gas or nuclear fuel"),
    ("What is the gravitational potential energy store?", "The energy an object has because of its position in a gravitational field"),
]

ATOMS_SECTION = [
    ("What is the relative charge of a proton?", "+1"),
    ("What is the relative charge of an electron?", "-1"),
    ("What is the relative charge of a neutron?", "0"),
    ("Where are protons and neutrons found in an atom?", "In the nucleus"),
    ("What is the atomic number of an element?", "The number of protons in an atom"),
    ("What is the mass number?", "The total number of protons and neutrons"),
    ("What are isotopes?", "Atoms of the same element with different numbers of neutrons"),
    ("Who discovered the electron?", "J.J. Thomson"),
    ("What did the alpha particle scattering experiment show?", "That the mass of an atom is concentrated in a small positive nucleus"),
    ("How many electrons fit in the first shell?", "2"),
]

ECOLOGY_QA = [
    ("What is a habitat?", "The place where an organism lives"),
    ("What is a population?", "All the organisms of one species living in a habitat"),
    ("What is a community?", "All the populations of different species living in a habitat"),
    ("What is an ecosystem?", "The interaction of a community of living organisms with the non-living parts of their environment"),
    ("What is a producer?", "An organism that makes its own food, usually by photosynthesis"),
    ("What is a biotic factor?", "A living factor that affects a community"),
    ("Give an example of an abiotic factor.", "Light intensity, temperature, moisture, soil pH or wind"),
    ("What does a quadrat measure?", "The number of organisms in a small sample area"),
    ("What is interdependence?", "When species in a community depend on each other for food, shelter, pollination or seed dispersal"),
    ("What do decomposers do?", "They break down dead organisms and return nutrients to the soil"),
]

SCANNED = [
    ("What is diffusion?", "The spreading out of particles from an area of higher concentration to an area of lower concentration"),
    ("What is osmosis?", "The diffusion of water through a partially permeable membrane"),
    ("What is active transport?", "The movement of substances against a concentration gradient using energy from respiration"),
    ("Where is the energy for active transport released?", "In respiration"),
    ("What is a partially permeable membrane?", "A membrane that lets some molecules through but not others"),
    ("Which process moves mineral ions into root hair cells?", "Active transport"),
    ("Give one factor that increases the rate of diffusion.", "A bigger concentration gradient, higher temperature or larger surface area"),
    ("What happens to an animal cell placed in pure water?", "It swells and may burst"),
]

PHOTO_KEYWORDS = [
    ("Conductor", "A material that allows electric current to flow through it easily"),
    ("Insulator", "A material that does not allow electric current to flow through it easily"),
    ("Current", "The rate of flow of electric charge"),
    ("Potential difference", "The energy transferred per unit charge between two points"),
    ("Resistance", "How much a component opposes the flow of current"),
    ("Series circuit", "A circuit with only one loop"),
    ("Parallel circuit", "A circuit with more than one loop"),
    ("Ammeter", "A meter that measures current, connected in series"),
]

CSV_QA = [
    ("What is the chemical symbol for sodium?", "Na"),
    ("What is the chemical symbol for iron?", "Fe"),
    ("What is the chemical symbol for potassium?", "K"),
    ("What is the formula of water?", "H2O"),
    ("What is the formula of carbon dioxide?", "CO2"),
    ("What is an element?", "A substance made of only one type of atom"),
    ("What is a compound?", "A substance made of two or more elements chemically joined"),
    ("What is a mixture?", "Two or more substances that are not chemically joined"),
    ("How can you separate an insoluble solid from a liquid?", "Filtration"),
    ("How can you separate a soluble solid from a solution?", "Crystallisation or evaporation"),
    ("What does chromatography separate?", "Mixtures of soluble substances, such as dyes in an ink"),
    ("How can you separate two liquids with different boiling points?", "Fractional distillation"),
    ("What is the formula of methane?", "CH4"),
    ("What is the formula of sodium chloride?", "NaCl"),
    ("What gas relights a glowing splint?", "Oxygen"),
]

XLSX_KEYWORDS = [
    ("Wavelength", "The distance from a point on one wave to the same point on the next wave"),
    ("Amplitude", "The maximum displacement of a point on a wave from its rest position"),
    ("Frequency", "The number of waves passing a point each second"),
    ("Period", "The time taken for one complete wave to pass a point"),
    ("Transverse wave", "A wave where the oscillations are perpendicular to the direction of energy transfer"),
    ("Longitudinal wave", "A wave where the oscillations are parallel to the direction of energy transfer"),
    ("Hertz", "The unit of frequency"),
    ("Reflection", "When a wave bounces off a surface"),
    ("Refraction", "When a wave changes direction as it passes from one medium into another"),
    ("Wave speed", "Frequency × wavelength"),
]

# Cloze: (sentence with the blank, the missing word(s)). Expected answer is the blank.
CLOZE = [
    ("Plants make glucose by ________.", "photosynthesis"),
    ("Photosynthesis takes place in the ________ of plant cells.", "chloroplasts"),
    ("The green pigment that absorbs light is called ________.", "chlorophyll"),
    ("The gas taken in by plants for photosynthesis is ________ ________.", "carbon dioxide"),
    ("The gas given out by photosynthesis is ________.", "oxygen"),
    ("Photosynthesis is an ________ reaction because it takes in energy.", "endothermic"),
    ("Water enters a plant through its ________ hair cells.", "root"),
    ("Glucose can be stored in plants as ________.", "starch"),
]

QUESTIONS_ONLY = [
    "What is an enzyme?",
    "What is the active site of an enzyme?",
    "What does amylase break down?",
    "What does protease break down into amino acids?",
    "Where is bile produced?",
    "What happens to an enzyme at a very high temperature?",
    "What is the optimum pH of most enzymes in the body?",
    "What does lipase break down?",
    "Where is amylase produced?",
    "Why is the small intestine good at absorbing food molecules?",
]

CLASS_LIST_DECK = [
    ("What is the speed of light in a vacuum?", "300 000 000 m/s"),
    ("Name the part of the electromagnetic spectrum with the longest wavelength.", "Radio waves"),
    ("Which electromagnetic waves are used to cook food?", "Microwaves"),
    ("Which electromagnetic waves are used in remote controls?", "Infrared"),
    ("Which electromagnetic waves can cause skin cancer?", "Ultraviolet"),
    ("Which electromagnetic waves are used to image broken bones?", "X-rays"),
]
CLASS_NAMES = ["Amelia Hart", "Jacob Singh", "Olivia Brennan", "Noah Kowalski", "Isla McCarthy", "Leo Adeyemi"]

EXPECTED = {}


def expect(name, pairs, shape, answers_missing=False, forbidden=None):
    EXPECTED[name] = {
        "shape": shape,
        "pairs": [{"q": q, "a": (None if answers_missing else a)} for q, a in pairs],
        "answers_missing": answers_missing,
        "forbidden": forbidden or [],
    }


# ── pptx ────────────────────────────────────────────────────────────────────
def pptx_next_slide():
    from pptx import Presentation
    from pptx.util import Pt, Inches
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Year 10 Forces — Retrieval Flashcards"
    title.placeholders[1].text = "Mr Badmus · Rainford High"
    for i, (q, a) in enumerate(FORCES_20, 1):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1)).text_frame
        tb.text = f"Question {i}"
        tb.paragraphs[0].runs[0].font.size = Pt(20)
        body = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(3)).text_frame
        body.word_wrap = True
        body.text = q
        body.paragraphs[0].runs[0].font.size = Pt(36)
        s2 = prs.slides.add_slide(blank)
        tb2 = s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1)).text_frame
        tb2.text = f"Answer {i}"
        body2 = s2.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(3)).text_frame
        body2.word_wrap = True
        body2.text = a
        body2.paragraphs[0].runs[0].font.size = Pt(32)
    name = "rainford_forces_next_slide.pptx"
    prs.save(os.path.join(OUT, name))
    expect(name, FORCES_20, "question on one slide, answer on the next")


def pptx_notes():
    from pptx import Presentation
    prs = Presentation()
    for q, a in CELLS_NOTES:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Cell structure — recall"
        s.placeholders[1].text = q
        s.notes_slide.notes_text_frame.text = f"Answer: {a}"
    name = "cells_speaker_notes.pptx"
    prs.save(os.path.join(OUT, name))
    expect(name, CELLS_NOTES, "answers in speaker notes")


def pptx_red():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    blank = prs.slide_layouts[5]
    s = prs.slides.add_slide(blank)
    s.shapes.title.text = "Acids and alkalis — answers in red"
    tf = s.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(9.2), Inches(5.5)).text_frame
    tf.word_wrap = True
    first = True
    for i, (q, a) in enumerate(ACIDS_RED, 1):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r1 = p.add_run()
        r1.text = f"{i}. {q} "
        r1.font.size = Pt(16)
        r2 = p.add_run()
        r2.text = a
        r2.font.size = Pt(16)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    name = "acids_red_answers.pptx"
    prs.save(os.path.join(OUT, name))
    expect(name, ACIDS_RED, "answers in red/bold on the same slide")


def pptx_class_list():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "10h/Ph1 — seating & groups"
    s.placeholders[1].text = "\n".join(CLASS_NAMES)
    for q, a in CLASS_LIST_DECK:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = "EM spectrum"
        rows = s.shapes.add_table(2, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(2)).table
        rows.cell(0, 0).text, rows.cell(0, 1).text = "Question", "Answer"
        rows.cell(1, 0).text, rows.cell(1, 1).text = q, a
    name = "em_spectrum_with_class_list.pptx"
    prs.save(os.path.join(OUT, name))
    expect(name, CLASS_LIST_DECK, "pupil names that must be ignored; one-row tables", forbidden=CLASS_NAMES)


# ── docx ────────────────────────────────────────────────────────────────────
def docx_table():
    import docx
    d = docx.Document()
    d.add_heading("Energy — key questions", 1)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Question", "Answer"
    for q, a in ENERGY_TABLE:
        c = t.add_row().cells
        c[0].text, c[1].text = q, a
    name = "energy_two_column_table.docx"
    d.save(os.path.join(OUT, name))
    expect(name, ENERGY_TABLE, "two-column table")


def docx_answer_section():
    import docx
    d = docx.Document()
    d.add_heading("Atomic structure — homework questions", 1)
    for i, (q, _) in enumerate(ATOMS_SECTION, 1):
        d.add_paragraph(f"{i}. {q}")
    d.add_page_break()
    d.add_heading("Answers", 2)
    for i, (_, a) in enumerate(ATOMS_SECTION, 1):
        d.add_paragraph(f"{i}. {a}")
    name = "atoms_answers_at_end.docx"
    d.save(os.path.join(OUT, name))
    expect(name, ATOMS_SECTION, "numbered questions with an answer section at the end")


# ── pdf ─────────────────────────────────────────────────────────────────────
def pdf_text():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    name = "ecology_qa_lines.pdf"
    c = canvas.Canvas(os.path.join(OUT, name), pagesize=A4)
    w, h = A4
    y = h - 60
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, "Ecology revision — Q and A")
    y -= 36
    for q, a in ECOLOGY_QA:
        for label, text, font in (("Q:", q, "Helvetica-Bold"), ("A:", a, "Helvetica")):
            c.setFont(font, 11)
            line = f"{label} {text}"
            while line:
                chunk = line[:95]
                if len(line) > 95 and " " in chunk:
                    chunk = chunk[:chunk.rindex(" ")]
                c.drawString(50, y, chunk)
                line = line[len(chunk):].lstrip()
                y -= 16
        y -= 10
        if y < 80:
            c.showPage()
            y = h - 60
    c.save()
    expect(name, ECOLOGY_QA, "Q:/A: lines in a text pdf")


def _sheet_image(title, rows, two_col, size=(1240, 1754)):
    img = Image.new("RGB", size, (250, 249, 244))
    d = ImageDraw.Draw(img)
    ft = ImageFont.truetype(FONT_B, 38)
    fb = ImageFont.truetype(FONT, 25)
    fbb = ImageFont.truetype(FONT_B, 25)
    d.text((80, 80), title, font=ft, fill=(20, 20, 20))
    y = 170
    for i, (left, right) in enumerate(rows, 1):
        if two_col:
            d.rectangle([70, y - 10, 1170, y + 120], outline=(60, 60, 60), width=2)
            d.line([420, y - 10, 420, y + 120], fill=(60, 60, 60), width=2)
            d.text((90, y + 35), left, font=fbb, fill=(15, 15, 15))
            words, line, yy = right.split(), "", y + 5
            for wd in words:
                if d.textlength(line + " " + wd, font=fb) > 720:
                    d.text((440, yy), line.strip(), font=fb, fill=(15, 15, 15)); yy += 34; line = ""
                line += " " + wd
            d.text((440, yy), line.strip(), font=fb, fill=(15, 15, 15))
            y += 130
        else:
            for label, text, f in ((f"{i}.", left, fbb), ("", right, fb)):
                words, line = text.split(), ""
                x0 = 80 if label else 130
                if label:
                    d.text((80, y), label, font=f, fill=(15, 15, 15)); x0 = 130
                for wd in words:
                    if d.textlength(line + " " + wd, font=f) > 1000:
                        d.text((x0, y), line.strip(), font=f, fill=(15, 15, 15)); y += 34; line = ""
                    line += " " + wd
                d.text((x0, y), line.strip(), font=f, fill=(15, 15, 15))
                y += 38
            y += 22
    return img


def pdf_scanned():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    img = _sheet_image("Transport in cells — questions and answers", SCANNED, two_col=False)
    img = img.rotate(0.8, expand=False, fillcolor=(240, 240, 235)).filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    img.convert("L").save(buf, "JPEG", quality=70)
    buf.seek(0)
    name = "transport_scanned.pdf"
    c = canvas.Canvas(os.path.join(OUT, name), pagesize=A4)
    w, h = A4
    c.drawImage(ImageReader(buf), 0, 0, width=w, height=h)
    c.save()
    expect(name, SCANNED, "scanned pdf (image only, question line then answer line)")


def photo():
    img = _sheet_image("Electricity key words", PHOTO_KEYWORDS, two_col=True)
    # A phone photo: on a desk, slightly rotated, perspective-ish shading, noise.
    canvasimg = Image.new("RGB", (1500, 2000), (120, 96, 70))
    img = img.rotate(-2.2, expand=True, fillcolor=(120, 96, 70))
    canvasimg.paste(img, (110, 90))
    shade = Image.new("L", canvasimg.size, 0)
    ds = ImageDraw.Draw(shade)
    for i in range(0, 2000, 4):
        ds.line([0, i, 1500, i], fill=int(40 * i / 2000))
    canvasimg = Image.composite(Image.new("RGB", canvasimg.size, (0, 0, 0)), canvasimg, shade)
    rnd = random.Random(351)
    px = canvasimg.load()
    for _ in range(60000):
        x, y = rnd.randrange(1500), rnd.randrange(2000)
        r, g, b = px[x, y]
        n = rnd.randint(-18, 18)
        px[x, y] = (max(0, min(255, r + n)), max(0, min(255, g + n)), max(0, min(255, b + n)))
    canvasimg = canvasimg.filter(ImageFilter.GaussianBlur(0.8))
    name = "electricity_keywords_photo.jpg"
    canvasimg.save(os.path.join(OUT, name), "JPEG", quality=78)
    expect(name, PHOTO_KEYWORDS, "photographed keyword → definition sheet")


# ── tabular / text ──────────────────────────────────────────────────────────
def csv_file():
    name = "chemistry_basics.csv"
    with open(os.path.join(OUT, name), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Question", "Answer"])
        for q, a in CSV_QA:
            w.writerow([q, a])
    expect(name, CSV_QA, "csv, header row")


def xlsx_file():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Waves keywords"
    ws.append(["Keyword", "Definition"])
    for k, v in XLSX_KEYWORDS:
        ws.append([k, v])
    name = "waves_keywords.xlsx"
    wb.save(os.path.join(OUT, name))
    expect(name, XLSX_KEYWORDS, "keyword → definition list in xlsx")


def cloze_txt():
    name = "photosynthesis_cloze.txt"
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        f.write("Photosynthesis — fill in the gaps\n\n")
        for i, (s, _) in enumerate(CLOZE, 1):
            f.write(f"{i}. {s}\n")
        f.write("\nWord bank: photosynthesis, chloroplasts, chlorophyll, carbon dioxide, oxygen, endothermic, root, starch\n")
    # The question wording is the model's to choose; scoring matches the answer
    # and requires the sentence's key content to survive in the question.
    expect(name, [(s.replace("________", "___"), a) for s, a in CLOZE], "cloze sentences with a word bank")


def questions_only_md():
    name = "enzymes_questions_only.md"
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        f.write("# Enzymes and digestion — retrieval questions\n\n")
        for i, q in enumerate(QUESTIONS_ONLY, 1):
            f.write(f"{i}. {q}\n")
    expect(name, [(q, None) for q in QUESTIONS_ONLY], "questions with no answers at all", answers_missing=True)


def main():
    os.makedirs(OUT, exist_ok=True)
    pptx_next_slide(); pptx_notes(); pptx_red(); pptx_class_list()
    docx_table(); docx_answer_section()
    pdf_text(); pdf_scanned(); photo()
    csv_file(); xlsx_file(); cloze_txt(); questions_only_md()
    with open(os.path.join(OUT, "expected.json"), "w", encoding="utf-8") as f:
        json.dump(EXPECTED, f, indent=2, ensure_ascii=False)
    total = sum(len(v["pairs"]) for v in EXPECTED.values())
    print(f"{len(EXPECTED)} fixtures, {total} expected pairs → {OUT}")


if __name__ == "__main__":
    main()
